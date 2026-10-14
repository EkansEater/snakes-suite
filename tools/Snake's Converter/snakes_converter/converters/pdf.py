"""Converting PDF files to images, text and Word documents."""

from __future__ import annotations

import re
import threading
from pathlib import Path

from ..core import ConversionError, Options, Progress, check_cancel, unique_path

PDF_OUT = ["docx", "txt", "png", "jpg"]

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _open(src: Path):
    import pypdfium2 as pdfium

    try:
        return pdfium.PdfDocument(str(src))
    except pdfium.PdfiumError as exc:
        if "password" in str(exc).lower():
            raise ConversionError("This PDF is password protected.") from None
        raise ConversionError(f"The PDF could not be read: {exc}") from None


def _page_text(page) -> str:
    textpage = page.get_textpage()
    try:
        return textpage.get_text_bounded().replace("\r\n", "\n").replace("\r", "\n")
    finally:
        textpage.close()


def convert(src: Path, dst: Path, target: str, opts: Options,
            progress: Progress, cancel: threading.Event) -> list[Path]:
    pdf = _open(src)
    try:
        count = len(pdf)
        if count == 0:
            raise ConversionError("The PDF has no pages.")
        if target in ("png", "jpg"):
            return _to_images(pdf, count, dst, target, opts, progress, cancel)

        pages = []
        for i in range(count):
            check_cancel(cancel)
            pages.append(_page_text(pdf[i]))
            progress((i + 1) / count)
        if not any(p.strip() for p in pages):
            raise ConversionError(
                "No text was found. This PDF is probably scanned images; convert it to PNG instead."
            )
        if target == "txt":
            dst.write_text("\n\n".join(pages), encoding="utf-8")
        else:
            _to_docx(pages, dst)
        return [dst]
    finally:
        pdf.close()


def _to_images(pdf, count: int, dst: Path, target: str, opts: Options,
               progress: Progress, cancel: threading.Event) -> list[Path]:
    digits = max(2, len(str(count)))
    outputs = []
    for i in range(count):
        check_cancel(cancel)
        if count == 1:
            path = dst
        else:
            path = dst.with_name(f"{dst.stem}_p{i + 1:0{digits}d}{dst.suffix}")
            if not opts.overwrite:
                path = unique_path(path)
        image = pdf[i].render(scale=opts.pdf_dpi / 72).to_pil().convert("RGB")
        if target == "jpg":
            image.save(path, "JPEG", quality=opts.image_quality, optimize=True)
        else:
            image.save(path, "PNG", optimize=True)
        outputs.append(path)
        progress((i + 1) / count)
    return outputs


def _to_docx(pages: list[str], dst: Path) -> None:
    """Text-only export: PDF layout, images and tables are not reconstructed."""
    import docx

    out = docx.Document()
    for index, text in enumerate(pages):
        if index:
            out.add_page_break()
        for paragraph in re.split(r"\n\s*\n", _CONTROL.sub("", text).strip()):
            if paragraph.strip():
                out.add_paragraph(paragraph)
    out.save(str(dst))
