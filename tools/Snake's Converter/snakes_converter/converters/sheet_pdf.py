"""Spreadsheets to PDF without LibreOffice: each sheet as a printed table, with its
number formats, bold, colours, fills, column widths and merged cells, the header row
repeated on every page, and landscape pages for wide sheets."""

from __future__ import annotations

import datetime
import html
import re
from pathlib import Path


def _fmt_number(v, fmt: str) -> str:
    """The common Excel number formats; anything else as the plain value."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime.datetime, datetime.date)):
        f = (fmt or "").lower()
        if "h" in f and ("mm" in f or ":" in f):
            return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime.datetime) else v.strftime("%Y-%m-%d")
        if "mmm" in f:
            return v.strftime("%b %-d, %Y") if "yy" in f else v.strftime("%b %-d")
        if f.startswith("m/") or f.startswith("mm/"):
            return v.strftime("%m/%d/%Y")
        if f.startswith("d/") or f.startswith("dd/"):
            return v.strftime("%d/%m/%Y")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if not isinstance(v, (int, float)):
        return "" if v is None else str(v)
    f = (fmt or "General").split(";")[0]
    if f == "General":
        return str(int(v)) if float(v).is_integer() and abs(v) < 1e15 else f"{v:.10g}"
    percent = "%" in f
    x = v * 100 if percent else v
    dec = 0
    m = re.search(r"0\.(0+)", f)
    if m:
        dec = len(m.group(1))
    s = f"{abs(x):,.{dec}f}" if "," in f else f"{abs(x):.{dec}f}"
    prefix = ""
    m = re.match(r'^(?:"([^"]*)"|\[\$([^-\]]*)[^\]]*\]|([$€£¥]))', f)
    if m:
        prefix = m.group(1) or m.group(2) or m.group(3) or ""
    suffix = ""
    m = re.search(r'"([^"]*)"\s*$', f)
    if m and not f.startswith('"'):
        suffix = m.group(1)
    out = f"{prefix}{s}{'%' if percent else ''}{suffix}"
    return f"-{out}" if x < 0 else out


def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) >= 6 and c.rgb not in ("00000000",):
            return "#" + c.rgb[-6:]
    except Exception:
        pass
    return None


def sheets_html(src: Path, ext: str, plain_sheets) -> list[tuple[str, str, int]]:
    """(name, html table, width in points) for each sheet."""
    out = []
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        from openpyxl.utils import get_column_letter

        book = openpyxl.load_workbook(src, data_only=True)
        for ws in book.worksheets:
            if ws.sheet_state != "visible":
                continue
            max_r, max_c = ws.max_row, ws.max_column
            # Trim empty edges.
            while max_r > 0 and all(ws.cell(max_r, c).value is None for c in range(1, max_c + 1)):
                max_r -= 1
            while max_c > 0 and all(ws.cell(r, max_c).value is None for r in range(1, max_r + 1)):
                max_c -= 1
            if not max_r or not max_c:
                continue
            spans = {}
            hidden_cells = set()
            for rng in ws.merged_cells.ranges:
                spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
                for r in range(rng.min_row, rng.max_row + 1):
                    for c in range(rng.min_col, rng.max_col + 1):
                        if (r, c) != (rng.min_row, rng.min_col):
                            hidden_cells.add((r, c))
            widths = []
            for c in range(1, max_c + 1):
                d = ws.column_dimensions.get(get_column_letter(c))
                hidden = d is not None and d.hidden
                w = (d.width if d is not None and d.width else 8.43) * 7 + 5  # characters -> pixels, as Excel shows them
                widths.append(0 if hidden else w)
            hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
            rows = []
            for r in range(1, max_r + 1):
                if r in hidden_rows:
                    continue
                cells = []
                for c in range(1, max_c + 1):
                    if (r, c) in hidden_cells or not widths[c - 1]:
                        continue
                    cell = ws.cell(r, c)
                    text = html.escape(_fmt_number(cell.value, cell.number_format)).replace("\n", "<br>")
                    st = []
                    font = cell.font
                    if font is not None:
                        if font.b:
                            text = f"<b>{text}</b>"
                        if font.i:
                            text = f"<i>{text}</i>"
                        col = _color(font.color)
                        if col and col.lower() != "#000000":
                            st.append(f"color:{col}")
                    fill = cell.fill
                    bg = _color(fill.fgColor) if fill is not None and fill.fill_type == "solid" else None
                    align = cell.alignment.horizontal if cell.alignment is not None else None
                    if not align:
                        align = "right" if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) else "left"
                    align = {"centerContinuous": "center", "general": "left"}.get(align, align)
                    rs, cs = spans.get((r, c), (1, 1))
                    attrs = f' align="{align}"' + (f' bgcolor="{bg}"' if bg else "") + (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
                    cells.append(f'<td{attrs} width="{int(widths[c - 1])}"' + (f' style="{";".join(st)}"' if st else "") + f">{text or '&nbsp;'}</td>")
                rows.append("<tr>" + "".join(cells) + "</tr>")
            head = rows[0] if ws.freeze_panes and ws.freeze_panes != "A1" or (ws.auto_filter and ws.auto_filter.ref) else ""
            body = rows[1:] if head else rows
            table = ('<table border="0" cellspacing="0" cellpadding="3">'
                     + (f"<thead>{head}</thead>" if head else "") + "<tbody>" + "".join(body) + "</tbody></table>")
            out.append((ws.title, table, int(sum(widths))))
        return out
    for name, rows in plain_sheets:
        if not rows:
            continue
        n = max(len(r) for r in rows)
        trs = []
        for i, row in enumerate(rows):
            cells = [html.escape("" if v is None else str(v)) for v in row] + [""] * (n - len(row))
            tag = "th" if i == 0 else "td"
            trs.append("<tr>" + "".join(f"<{tag}>{c or '&nbsp;'}</{tag}>" for c in cells) + "</tr>")
        table = ('<table border="1" cellspacing="0" cellpadding="3" style="border-collapse:collapse;border-color:#bbbbbb">'
                 f"<thead>{trs[0]}</thead><tbody>{''.join(trs[1:])}</tbody></table>")
        out.append((name, table, n * 70))
    return out


def write_pdf(parts, dst: Path, title: str) -> None:
    from PySide6.QtCore import QMarginsF, QRectF, QSizeF, Qt
    from PySide6.QtGui import QColor, QFont, QPageLayout, QPageSize, QPainter, QPdfWriter, QTextDocument

    if not parts:
        from ..core import ConversionError

        raise ConversionError("The spreadsheet is empty.")
    wide = max(w for _, _, w in parts) > 720
    writer = QPdfWriter(str(dst))
    writer.setResolution(300)
    size = QPageSize(QPageSize.PageSizeId.A4)
    orient = QPageLayout.Orientation.Landscape if wide else QPageLayout.Orientation.Portrait
    writer.setPageLayout(QPageLayout(size, orient, QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter))
    writer.setTitle(title)
    dpmm = writer.resolution() / 25.4
    full = writer.pageLayout().fullRectPixels(writer.resolution())
    m = 12 * dpmm
    body = QRectF(m, m + 6 * dpmm, full.width() - 2 * m, full.height() - 2 * m - 12 * dpmm)
    painter = QPainter(writer)
    first = True
    try:
        for name, table, width in parts:
            doc = QTextDocument()
            doc.setDocumentMargin(0)
            font = QFont("Helvetica", 9)
            doc.setDefaultFont(font)
            doc.documentLayout().setPaintDevice(writer)
            doc.setPageSize(body.size())
            # Tables wider than the page are shrunk to fit its width.
            # Widths are in screen pixels (1/96 inch), as Qt reads HTML.
            scale = min(1.0, (body.width() / (writer.resolution() / 96)) / max(1, width))
            sized = re.sub(r'<table border="(\d)"', lambda mm: f'<table width="{int(width)}" border="{mm.group(1)}"', table, count=1)
            doc.setDefaultStyleSheet("td { white-space: nowrap; }")
            doc.setHtml(sized)
            if scale < 1:
                doc.setPageSize(QSizeF(body.width() / scale, body.height() / scale))
            pages = max(1, doc.pageCount())
            for i in range(pages):
                if not first:
                    writer.newPage()
                first = False
                painter.save()
                painter.translate(body.left(), body.top())
                painter.scale(scale, scale)
                ph = body.height() / scale
                painter.translate(0, -i * ph)
                doc.drawContents(painter, QRectF(0, i * ph, body.width() / scale, ph - 2))
                painter.restore()
                painter.setFont(QFont("Helvetica", 8))
                painter.setPen(QColor("#555555"))
                painter.drawText(QRectF(body.left(), 3 * dpmm, body.width(), 8 * dpmm), int(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter), name)
                painter.drawText(QRectF(body.left(), full.height() - 10 * dpmm, body.width(), 8 * dpmm), int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter), f"{i + 1} / {pages}")
    finally:
        painter.end()
