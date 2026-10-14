"""Spreadsheet and tabular data conversion (XLSX, CSV, TSV, JSON)."""

from __future__ import annotations

import csv
import datetime
import io
import json
import re
import threading
from pathlib import Path

from ..core import ConversionError, Options, Progress, check_cancel, read_text, unique_path
from . import office

SHEET_IN = {"xlsx", "xlsm", "xls", "ods", "csv", "tsv", "json"}

_ORDER = ["xlsx", "csv", "pdf", "ods", "tsv", "json"]
_BUILTIN_IN = {"xlsx", "xlsm", "csv", "tsv", "json"}
_BUILTIN_OUT = {"xlsx", "csv", "tsv", "json"}
_OFFICE_IN = {"xlsx", "xlsm", "xls", "ods", "csv"}
_OFFICE_OUT = {"xlsx", "csv", "pdf", "ods"}

Sheets = list[tuple[str, list[list]]]  # (sheet name, rows)


def targets(ext: str) -> list[str]:
    possible = set()
    if ext in _BUILTIN_IN:
        possible |= _BUILTIN_OUT
    if office.available() and ext in _OFFICE_IN:
        possible |= _OFFICE_OUT
    possible.discard("xlsx" if ext == "xlsm" else ext)
    return [t for t in _ORDER if t in possible]


def convert(src: Path, dst: Path, target: str, opts: Options,
            progress: Progress, cancel: threading.Event) -> list[Path]:
    ext = src.suffix.lower().lstrip(".")
    if ext in _BUILTIN_IN and target in _BUILTIN_OUT:
        try:
            sheets = _read(src, ext)
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"The file could not be read: {exc}") from None
        check_cancel(cancel)
        return _write(sheets, dst, target, opts)
    return office.convert(src, dst, target, progress, cancel)


# --- reading ---------------------------------------------------------------

def _read(src: Path, ext: str) -> Sheets:
    if ext in ("xlsx", "xlsm"):
        import openpyxl

        book = openpyxl.load_workbook(src, read_only=True, data_only=True)
        try:
            return [(ws.title, [list(row) for row in ws.iter_rows(values_only=True)])
                    for ws in book.worksheets]
        finally:
            book.close()
    text = read_text(src)
    if ext == "json":
        return [(src.stem, _rows_from_json(text))]
    delimiter = "\t" if ext == "tsv" else _sniff(text)
    return [(src.stem, list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)))]


def _sniff(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _rows_from_json(text: str) -> list[list]:
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ConversionError(f"The JSON file is not valid: {exc}") from None
    if isinstance(data, dict):
        # Accept {"items": [...]} style wrappers around the actual table.
        lists = [v for v in data.values() if isinstance(v, list)]
        data = lists[0] if len(lists) == 1 else [data]
    if not isinstance(data, list):
        raise ConversionError("The JSON file must contain a list of records.")
    if all(isinstance(item, dict) for item in data):
        header = list(dict.fromkeys(key for item in data for key in item))
        return [header] + [[_cell(item.get(key)) for key in header] for item in data]
    return [[_cell(v) for v in item] if isinstance(item, list) else [_cell(item)] for item in data]


def _cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


# --- writing ---------------------------------------------------------------

def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _typed(value):
    """Turn CSV strings back into numbers where that is unambiguous."""
    if not isinstance(value, str):
        return value
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})", value):
        return int(value)
    if re.fullmatch(r"-?(0|[1-9]\d*)\.\d+", value) and len(value) <= 16:
        return float(value)
    return value


def _json_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return value


def _write(sheets: Sheets, dst: Path, target: str, opts: Options) -> list[Path]:
    if target == "xlsx":
        import openpyxl

        book = openpyxl.Workbook(write_only=True)
        for name, rows in sheets:
            ws = book.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet")
            for row in rows:
                ws.append([_typed(v) for v in row])
        book.save(dst)
        return [dst]

    # One output file per sheet; a single sheet keeps the plain file name.
    outputs = []
    for name, rows in sheets:
        path = dst
        if len(sheets) > 1:
            safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip() or "Sheet"
            path = dst.with_name(f"{dst.stem}_{safe}{dst.suffix}")
            if not opts.overwrite:
                path = unique_path(path)
        if target == "json":
            header = [_text(h) or f"column_{i + 1}" for i, h in enumerate(rows[0])] if rows else []
            records = [
                {key: _json_value(_typed(row[i]) if i < len(row) else None)
                 for i, key in enumerate(header)}
                for row in rows[1:]
            ]
            path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
        else:
            with open(path, "w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(fh, delimiter="\t" if target == "tsv" else ",")
                writer.writerows([_text(v) for v in row] for row in rows)
        outputs.append(path)
    return outputs
